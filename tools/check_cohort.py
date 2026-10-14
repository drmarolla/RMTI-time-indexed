# Copyright 2026 Cesar Marolla
# Licensed under the Apache License, Version 2.0 (see LICENSE).
"""
check_cohort.py -- the Python cohort, the disaster workbook and the disaster
sandbox must describe the same worked cases.

    python tools/check_cohort.py            # needs: pip install openpyxl

Source of truth: COHORT in rmti-python/disaster_climate/rmti_ssa.py (the values
reported in Tables 4-6 and the January 2025 wildfire cases). For every case it
compares, field by field:
  * the workbook's RMTI Calculator sheet: inputs N, L, E, V, rho now, rho after,
    C, and the calculated R now, R after, EAL now, EAL after (cached values; open
    and recalculate the workbook before running this after any edit);
  * the sandbox's example list in disaster/index.html: the same inputs.
Exit status 0 = all three agree; 1 = at least one difference; 2 = cannot run.
"""
import os
import re
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "rmti-python"))
try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is required: pip install openpyxl")
    sys.exit(2)
from core.rmti_core import score                                    # noqa: E402
from disaster_climate.rmti_ssa import COHORT                         # noqa: E402

TOL = 1e-9
problems = 0


def report(msg):
    global problems
    problems += 1
    print("  DIFFERENCE:", msg)


# --------------------------------------------------------------- workbook
ws = load_workbook(os.path.join(ROOT, "calculators", "RMTI_Calculator_Disaster_v2.3.0.xlsx"),
                   data_only=True)["RMTI Calculator"]
rows = {str(ws.cell(r, 1).value).replace("★", "").strip(): r for r in range(5, 18)
        if ws.cell(r, 1).value}
checked = 0
for city, m in COHORT.items():
    if city not in rows:
        report(f"workbook has no row for {city}")
        continue
    r = rows[city]
    want = {5: m["N"], 6: m["L"], 8: m["E_sel"], 9: m["V"], 10: m["rho_now"],
            11: m["rho_after"], 12: m["C"]}
    names = {5: "N", 6: "L", 8: "E used", 9: "V", 10: "rho now", 11: "rho after", 12: "C"}
    for col, val in want.items():
        got = ws.cell(r, col).value
        if not isinstance(got, (int, float)) or abs(got - val) > TOL:
            report(f"workbook {city}: {names[col]} = {got!r}, Python cohort has {val!r}")
    s0 = score(m["L"], m["E_sel"], m["V"], m["rho_now"], m["C"])
    s1 = score(m["L"], m["E_sel"], m["V"], m["rho_after"], m["C"])
    for col, val, nm in ((15, s0["R_raw"], "R now"), (17, s1["R_raw"], "R after"),
                         (20, s0["EAL_raw"], "EAL now"), (21, s1["EAL_raw"], "EAL after")):
        got = ws.cell(r, col).value
        if not isinstance(got, (int, float)) or abs(got - val) > TOL:
            report(f"workbook {city}: {nm} = {got!r}, Python gives {val!r}")
    checked += 1
print(f"workbook: {checked} of {len(COHORT)} cohort cases compared")

# --------------------------------------------------------------- sandbox
html = open(os.path.join(ROOT, "disaster", "index.html"), encoding="utf-8").read()
ex = {}
for mm in re.finditer(r"\{name:'([^']+)',\s*N:(\d+),\s*L:([\d.]+),\s*E:([\d.]+),\s*V:([\d.]+),"
                      r"\s*rhoNow:([\d.]+),\s*rhoAfter:([\d.]+),\s*C:(\d)\}", html):
    ex[mm.group(1).replace("★", "").strip()] = [float(x) for x in mm.groups()[1:]]
sb = 0
for city, m in COHORT.items():
    if city not in ex:
        continue                         # the sandbox may show a subset of cases
    got = ex[city]
    want = [m["N"], m["L"], m["E_sel"], m["V"], m["rho_now"], m["rho_after"], m["C"]]
    for nm, g, w in zip(("N", "L", "E", "V", "rho now", "rho after", "C"), got, want):
        if abs(g - w) > TOL:
            report(f"sandbox {city}: {nm} = {g}, Python cohort has {w}")
    sb += 1
if sb == 0:
    report("no worked-example rows could be read from disaster/index.html")
print(f"sandbox:  {sb} of {len(COHORT)} cohort cases compared")

print("OK: Python cohort, workbook and sandbox agree" if not problems
      else f"{problems} difference(s) found")
sys.exit(1 if problems else 0)
