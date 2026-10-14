# Copyright 2026 Cesar Marolla
# Licensed under the Apache License, Version 2.0 (see LICENSE).
"""
check_workbooks.py -- verify that each calculator's Trajectory sheet agrees with
the Python implementation (dynamic/rmti_dynamic.py).

    python tools/check_workbooks.py            # needs: pip install openpyxl

For each workbook it reads the sheet's OWN parameters (B5:B13), inputs (L, E, C,
rho at t=0) and observations, runs the Python code on exactly those, and compares
level, slope, V*, rho and R for every cycle of both example blocks, plus the tier
number. Cached values are used, so open and save the workbook in Excel or
LibreOffice (a full recalculation) before running this after any edit.
"""
import os
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "rmti-python"))
try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is required: pip install openpyxl")
    sys.exit(2)
from dynamic.rmti_dynamic import TrajectoryParams, run_trajectory   # noqa: E402

TOL = 1e-12
BLOCKS = [(16, 22), (32, 38)]          # (first input row, first cycle row) per example
failed = 0
for fn in ("RMTI_HCC_Calculator_v2.3.0.xlsx", "RMTI_Calculator_Disaster_v2.3.0.xlsx"):
    ws = load_workbook(os.path.join(ROOT, "calculators", fn), data_only=True)["Trajectory"]
    p = TrajectoryParams(v_lo=ws["B5"].value, v_hi=ws["B6"].value, y_max=ws["B7"].value,
                         alpha=ws["B8"].value, beta=ws["B9"].value, look_ahead=ws["B10"].value,
                         lam=ws["B11"].value, rho_ceiling=ws["B12"].value,
                         filter_on=bool(ws["B13"].value))
    worst, n = 0.0, 0
    for r_in, r0 in BLOCKS:
        L, E, C, rho0 = (ws.cell(r_in + i, 2).value for i in range(4))
        obs, dlv, r = [], [], r0
        while ws.cell(r, 1).value is not None and ws.cell(r, 2).value is not None:
            obs.append(ws.cell(r, 2).value)
            dlv.append(ws.cell(r, 3).value)
            r += 1
        rows = run_trajectory(L, E, int(C), rho0, obs, dlv, p)
        for i, row in enumerate(rows):
            rr = r0 + i
            for col, key in ((4, "level"), (5, "slope"), (6, "V"), (7, "rho"), (8, "R")):
                got = ws.cell(rr, col).value
                if not isinstance(got, (int, float)):
                    print(f"  {fn} row {rr}: no calculated value found ({got!r}). The workbook was saved "
                          "without recalculation; open it in Excel/LibreOffice, recalculate, save, and rerun.")
                    sys.exit(2)
                diff = abs(got - row[key])
                worst, n = max(worst, diff), n + 1
                if diff > TOL:
                    print(f"  MISMATCH {fn} row {rr} {key}: workbook {got} vs python {row[key]}")
                    failed += 1
            tier_xl = str(ws.cell(rr, 9).value)[0]
            if tier_xl != row["tier"].split()[1]:
                print(f"  TIER MISMATCH {fn} row {rr}: {tier_xl} vs {row['tier']}")
                failed += 1
    print(f"{fn}: {n} values compared, max |difference| = {worst:.2e}")
print("OK: workbooks agree with the Python code" if not failed else f"{failed} mismatch(es)")
sys.exit(1 if failed else 0)
