# Copyright 2026 Cesar Marolla
# Licensed under the Apache License, Version 2.0 (see LICENSE).
"""
check_release.py -- pre-release consistency gate for this repository.

    python tools/check_release.py             # normal run (DOI placeholders = warnings)
    python tools/check_release.py --release   # strict: DOI placeholders must be filled

Exit status 0 = every check passed. Run it before tagging a release, and
again after Zenodo has issued the DOI and the placeholders are filled.
Standard library only (PyYAML and openpyxl are used if installed).
"""
import json
import os
import re
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
COMPONENTS = ["rmti-python", "calculators", "oncology", "disaster"]
failures, warnings, doi_todo = [], [], []


def p(*a):
    return os.path.join(ROOT, *a)


def read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


def fail(msg):
    failures.append(msg)


def warn(msg):
    warnings.append(msg)


# ------------------------------------------------------------------- folder layout
# Runs first. Guards against a flattened upload (files copied out of their folders):
# if the layout is wrong nothing else can be checked, so stop with a clear message.
import glob
REQUIRED = [
    ".github/workflows/tests.yml", ".gitignore", ".zenodo.json",
    "rmti-python/pyproject.toml", "rmti-python/reproduce.py",
    "rmti-python/core/__init__.py", "rmti-python/core/rmti_core.py",
    "rmti-python/dynamic/__init__.py", "rmti-python/dynamic/rmti_dynamic.py",
    "rmti-python/oncology/__init__.py", "rmti-python/oncology/rmti_hcc.py",
    "rmti-python/disaster_climate/__init__.py", "rmti-python/disaster_climate/rmti_ssa.py",
    "rmti-python/tests/__init__.py", "rmti-python/tests/test_rmti.py", "rmti-python/tests/test_dynamic.py",
    "oncology/index.html", "disaster/index.html", "index.html",
    "docs/TRAJECTORY_GUIDE.md", "docs/PAPER_SUPPORT.md", "docs/reproduce_expected_output.txt",
    "tools/check_release.py", "tools/check_workbooks.py", "tools/check_cohort.py", "tools/set_doi.py",
]
for rel in REQUIRED:
    if not os.path.exists(p(*rel.split("/"))):
        fail(f"required file missing or in the wrong folder: {rel}")
for pattern in ("calculators/RMTI_HCC_Calculator_v*.xlsx", "calculators/RMTI_Calculator_Disaster_v*.xlsx",
                "docs/RMTI_Calculator_Explainer_v*.docx"):
    if not glob.glob(p(*pattern.split("/"))):
        fail(f"required file missing or in the wrong folder: {pattern}")
for fn in os.listdir(ROOT):
    full = os.path.join(ROOT, fn)
    if os.path.isfile(full) and (fn.endswith((".py", ".xlsx", ".docx", ".yml", ".toml", ".txt"))
                                 or re.search(r" \(\d+\)", fn) or fn in ("download", "zenodo.json")):
        fail(f"stray file at the repository root (flattened upload?): {fn}")
if failures:
    print("RMTI release check: the folder layout is wrong, so nothing else was checked.")
    for f in failures:
        print("  FAIL:   ", f)
    print("\nRestore the layout from the release zip (see UPLOAD_INSTRUCTIONS.md, step 1) and rerun.")
    sys.exit(1)

# --------------------------------------------------------------------- version
core_src = read(p("rmti-python", "core", "rmti_core.py"))
m = re.search(r'__version__\s*=\s*"([^"]+)"', core_src)
VERSION = m.group(1) if m else None
if not VERSION:
    fail("core/rmti_core.py: __version__ not found")
    VERSION = "?"


def expect_version(label, found):
    if found != VERSION:
        fail(f"version mismatch in {label}: {found!r} (core says {VERSION!r})")


for mod in ("oncology/rmti_hcc.py", "disaster_climate/rmti_ssa.py"):
    mm = re.search(r'__version__\s*=\s*"([^"]+)"', read(p("rmti-python", *mod.split("/"))))
    expect_version(mod, mm.group(1) if mm else None)
mm = re.search(r'^version\s*=\s*"([^"]+)"', read(p("rmti-python", "pyproject.toml")), re.M)
expect_version("pyproject.toml", mm.group(1) if mm else None)

cff_paths = ["CITATION.cff"] + [f"{c}/CITATION.cff" for c in COMPONENTS]
for cp in cff_paths:
    if not os.path.exists(p(*cp.split("/"))):
        fail(f"missing {cp}")
        continue
    t = read(p(*cp.split("/")))
    mm = re.search(r"^version:\s*([^\s#]+)", t, re.M)
    expect_version(cp, mm.group(1).strip('"') if mm else None)
    if "license: Apache-2.0" not in t:
        fail(f"{cp}: licence is not Apache-2.0")
try:
    import yaml  # optional
    for cp in cff_paths:
        yaml.safe_load(read(p(*cp.split("/"))))
except ImportError:
    pass
except Exception as e:                                               # noqa: BLE001
    fail(f"CITATION.cff is not valid YAML: {e}")

# ---------------------------------------------------------------- .zenodo.json
try:
    z = json.loads(read(p(".zenodo.json")))
    for k in ("title", "description", "creators", "license", "upload_type", "access_right"):
        if k not in z:
            fail(f".zenodo.json: missing '{k}'")
    expect_version(".zenodo.json", z.get("version"))
    mt = re.search(r'^title:\s*"([^"]+)"', read(p("CITATION.cff")), re.M)
    if not mt or mt.group(1) != z.get("title"):
        fail("root CITATION.cff title differs from the .zenodo.json title "
             f"({mt.group(1) if mt else None!r} vs {z.get('title')!r})")
    if z.get("license") != "Apache-2.0":
        fail(".zenodo.json: license is not Apache-2.0")
    if z.get("upload_type") != "software":
        fail(".zenodo.json: upload_type is not 'software'")
    for c in z.get("creators", []):
        if not c.get("name"):
            fail(".zenodo.json: creator without a name")
except Exception as e:                                                   # noqa: BLE001
    fail(f".zenodo.json unreadable: {e}")

# ------------------------------------------------------------------- sandboxes
for sb in ("oncology", "disaster"):
    html = read(p(sb, "index.html"))
    mm = re.search(r"RMTI_RELEASE\s*=\s*\{[^}]*?version:\s*'([^']+)'", html, re.S)
    expect_version(f"{sb}/index.html RMTI_RELEASE", mm.group(1) if mm else None)
    if "drmarolla/RMTI-time-indexed" not in html:
        fail(f"{sb}/index.html: repo URL is not the time-indexed repository")
    for ph in re.findall(r"TODO[^\n]*", html):
        doi_todo.append(f"{sb}/index.html")

# ------------------------------------------------------------------ changelogs
for cl in ("CHANGELOG.md", "rmti-python/CHANGELOG.md"):
    mm = re.search(r"^## \[([^\]]+)\]", read(p(*cl.split("/"))), re.M)
    expect_version(cl + " (top entry)", mm.group(1) if mm else None)

# -------------------------------------------------------------------- workbooks
try:
    from openpyxl import load_workbook
    cover = load_workbook(p("calculators", "RMTI_HCC_Calculator_v2.3.0.xlsx"))["Cover"]
    text = " ".join(str(c.value) for r in cover.iter_rows() for c in r if c.value)
    if f"v{VERSION}" not in text:
        fail(f"HCC workbook Cover does not mention v{VERSION}")
    for fn in os.listdir(p("calculators")):
        if fn.endswith(".xlsx"):
            wb = load_workbook(p("calculators", fn))
            for ws in wb:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and re.search(
                                r"proprietary|not for redistribution", c.value, re.I):
                            fail(f"{fn}!{ws.title}!{c.coordinate}: contradicts Apache-2.0")
except ImportError:
    warn("openpyxl not installed: workbook checks skipped")


# ---------------------------------------------------------- licences and notices
root_lic = read(p("LICENSE"))
if "Apache License" not in root_lic.splitlines()[1] + root_lic.splitlines()[0]:
    fail("root LICENSE does not look like the Apache License 2.0")
for c in COMPONENTS:
    for fn in ("LICENSE", "NOTICE", "CITATION.cff", "DISCLAIMER.md", "README.md"):
        if not os.path.exists(p(c, fn)):
            fail(f"{c}/{fn} is missing")
    if os.path.exists(p(c, "LICENSE")) and read(p(c, "LICENSE")) != root_lic:
        fail(f"{c}/LICENSE differs from the root LICENSE")
for fn in ("NOTICE", "CITATION.cff", ".zenodo.json", "DISCLAIMER.md", "README.md",
           "CHANGELOG.md", "index.html", ".gitignore", "UPLOAD_INSTRUCTIONS.md"):
    if not os.path.exists(p(fn)):
        fail(f"root {fn} is missing")
dis = [read(p("DISCLAIMER.md"))] + [read(p(c, "DISCLAIMER.md")) for c in COMPONENTS
                                    if os.path.exists(p(c, "DISCLAIMER.md"))]
if len(set(dis)) != 1:
    fail("DISCLAIMER.md copies are not identical")

# ------------------------------------------------------- text hygiene / old names
SCAN_EXT = (".md", ".py", ".html", ".cff", ".json", ".yml", ".toml")
for dirpath, dirs, files in os.walk(ROOT):
    dirs[:] = [d for d in dirs if d not in (".git", "__pycache__")]
    for fn in files:
        full = os.path.join(dirpath, fn)
        rel = os.path.relpath(full, ROOT).replace(os.sep, "/")
        if fn.endswith(".pyc"):
            fail(f"stray bytecode file: {rel}")
        if fn.endswith(SCAN_EXT) or fn in ("NOTICE", "LICENSE"):
            if rel == "tools/check_release.py":
                continue
            t = read(full)
            if re.search(r"proprietary|not for redistribution", t, re.I) and fn != "CHANGELOG.md":
                fail(f"{rel}: contains a non-Apache licence statement")
            if re.search(r"(?<!rmti-)(oncology|disaster-science)-sandbox", t) and fn != "CHANGELOG.md":
                fail(f"{rel}: refers to an old sandbox folder name")
            if fn in ("NOTICE", "CITATION.cff") and "not yet" in t.lower():
                doi_todo.append(rel)
            if fn == "CITATION.cff" and re.search(r"^# doi:", t, re.M):
                doi_todo.append(rel)

if doi_todo:
    files = sorted(set(doi_todo))
    msg = f"DOI placeholders still present in {len(files)} file(s): " + ", ".join(files)
    (fail if "--release" in sys.argv else warn)(msg)

# ------------------------------------------------------------------- html links
index = read(p("index.html"))
for href in re.findall(r'href="([^"#:]+)"', index):
    target = p(*href.rstrip("/").split("/"))
    if not (os.path.exists(target) or os.path.exists(os.path.join(target, "index.html"))):
        fail(f"index.html links to a missing path: {href}")

# --------------------------------------------------------------------- web upload
n_files = sum(len(f) for _, d, f in os.walk(ROOT) if ".git" not in _.split(os.sep))
if n_files >= 100:
    warn(f"{n_files} files: GitHub's browser uploader accepts at most 100 per batch")
big = [f for d, _, fs in os.walk(ROOT) for f in fs
       if os.path.getsize(os.path.join(d, f)) > 25 * 1024 * 1024]
if big:
    fail(f"files over 25 MB (browser upload limit): {big}")

# ----------------------------------------------------------------------- report
print(f"RMTI release check -- version {VERSION} -- {n_files} files")
for w in warnings:
    print("  warning:", w)
for f in failures:
    print("  FAIL:   ", f)
if failures:
    print(f"\n{len(failures)} check(s) failed.")
    sys.exit(1)
print("\nAll checks passed" + (" (strict)." if "--release" in sys.argv else
                               f" ({len(warnings)} placeholder warning(s) expected before the DOI exists)."))
